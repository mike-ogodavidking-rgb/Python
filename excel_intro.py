from openpyxl import Workbook

wb = Workbook()
ws = wb.active
ws.title = "Scores"

ws["A1"] = "Name"
ws["B1"] = "Score"
ws.append(["Ada", 72])
ws.append(["Chi", 81])
ws.append(["Wisdom", 88])
ws["A5"] = "Total"
ws["B5"] = "=SUM(B2:B4)"

from openpyxl.styles import Font, PatternFill

for cell in ws[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="1E7B34", end_color="1E7B34", fill_type="solid")

for cell in ws["B"]:
    if cell.row != 1:
        cell.number_format = "0.00"

for cell in ws[5]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="1E7B34", end_color="1E7B34", fill_type="solid")

ws.column_dimensions["A"].width = 20
ws.column_dimensions["B"].width = 10
ws.freeze_panes = "A2"

wb.save("practice.xlsx")