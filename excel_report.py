import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

def csv_to_report(input_file, output_file):
    df = pd.read_csv(input_file)

    for column in df.columns:
        if pd.api.types.is_string_dtype(df[column]):
            df[column] = df[column].str.strip().str.title()

    df = df.drop_duplicates()
    df = df.dropna()
    df.columns = df.columns.str.title()

    df.to_excel(output_file, index=False, sheet_name="Report")

    wb = load_workbook(output_file)
    ws = wb["Report"]

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1E7B34", end_color="1E7B34", fill_type="solid")

    numeric_columns = df.select_dtypes(include="number").columns 
    last = ws.max_row

    ws.cell(row=last + 1, column=1, value="Average")
    ws.cell(row=last +1, column=1).font = Font(bold=True)

    for name in numeric_columns:
        position = df.columns.get_loc(name) + 1
        letter = get_column_letter(position)

        ws.cell(row=last + 1, column=position, value=f"=AVERAGE({letter}2:{letter}{last})")
        ws.cell(row=last + 1, column=position).number_format = "0.00"
        ws.cell(row=last +1, column=position).font = Font(bold=True)
        

    for column in ws.columns:
        longest = 0
        for cell in column:
            if cell.value is not None and not str(cell.value).startswith("="):
                longest = max(longest, len(str(cell.value)))
        ws.column_dimensions[column[0].column_letter].width = longest + 2

    ws.freeze_panes = "A2"
    wb.save(output_file)
    return output_file

while True:
    filename = input("Enter the CSV filename: ").strip()
    if not filename.endswith(".csv"):
        print("Please enter a valid CSV filename.")
        continue

    output_name = filename[:-4] + "_report.xlsx"

    try:
        saved = csv_to_report(filename, output_name)
        print(f"Report generated and saved as {saved}")
    except FileNotFoundError:
        print("File not found. Please check the filename and try again.")
    except PermissionError:
        print("Permission denied. Please close the Excel file if it's open and try again.")

    repeat = input("Do you want to generate another report? (yes/no): ").strip().lower()
    if repeat != "yes":
        break